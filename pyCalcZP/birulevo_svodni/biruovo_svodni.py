#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import os
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import column_index_from_string

import finwizard_sdk as fw

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="4F81BD")
header_align = Alignment(horizontal="center", vertical="center")
thin = Side(border_style="thin", color="000000")
header_border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(cell):
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align


columns_align = Alignment(horizontal="center", vertical="center")


def style_columns(cell):
    cell.alignment = columns_align


headers = [
    "Дата", "Смена", "Выручка за смену", "Услуги/день",
    "Услуги ночь", "Услуги всего", "З/п врачей/всего", "З/п 10%",
    "Ритуал", "Митрохина", "Ратолог",
    "Дерматолог", "Стоматолог", "Белозор", "Онколог", "Кардиолог",
    "Чеки день", "Чеки ночь", "Чеки аптека", "Аптека", "Неовет", "Веттест"
]

column_widths = {
    1: 12, 2: 70, 3: 35, 4: 15, 5: 15, 6: 15, 7: 15, 8: 15, 9: 15,
    10: 15, 11: 15, 12: 15, 13: 15, 14: 15, 15: 15, 16: 15, 17: 15,
    18: 15, 19: 15, 20: 15, 21: 15, 22: 15
}


def print_values(cell_address, col_out, input_ws, output_ws, row_out, filename):
    try:
        row_in = int(''.join(filter(str.isdigit, cell_address)))
        col_in = column_index_from_string(''.join(filter(str.isalpha, cell_address)))
        value = input_ws.cell(row=row_in, column=col_in).value
        if value is None:
            return
        value = float(value)
        cell = output_ws.cell(row=row_out, column=col_out, value=value)
        style_columns(cell)
    except Exception as e:
        print(f"[ОШИБКА] Не удалось прочитать {cell_address} в файле {filename}: {e}")


@fw.entrypoint
def run(params):
    file_names = params.get("xlsxFiles", [])
    output_path = params.get("outputFile")

    if not file_names:
        raise fw.PluginError("Список входных файлов пуст.")

    output_wb = Workbook()
    output_ws = output_wb.active
    output_ws.title = "БирюлевоОтчет"

    for col, title in enumerate(headers, start=1):
        style_header(output_ws.cell(row=1, column=col, value=title))

    for col, width in column_widths.items():
        output_ws.column_dimensions[chr(64 + col)].width = width

    rows_indent = 2
    total = len(file_names)

    for idx, filename in enumerate(file_names, start=1):
        fw.core.update_progress(int(idx / total * 100), f"Обработка {os.path.basename(filename)}")

        try:
            input_wb = openpyxl.load_workbook(filename, data_only=True)
        except Exception as e:
            print(f"[ОШИБКА] Не удалось открыть файл {filename}: {e}")
            continue
        input_ws = input_wb.active

        try:
            value = input_ws["E1"].value
            value = str(value).rstrip(".")
            cell = output_ws.cell(row=rows_indent, column=1, value=value)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        except Exception as e:
            print(f"[ОШИБКА] Не удалось прочитать E1 в файле {filename}: {e}")
            continue

        allnames_list = []
        try:
            for row in range(6, 8):
                for col in range(1, 4):
                    value = input_ws.cell(row=row, column=col).value
                    if value and str(value).strip():
                        allnames_list.append(str(value).strip())
            allnames = ", ".join(allnames_list)
            cell = output_ws.cell(row=rows_indent, column=2, value=allnames)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        except Exception as e:
            print(f"[ОШИБКА] Не удалось собрать имена в файле {filename}: {e}")
            continue

        print_values("D13", 3, input_ws, output_ws, rows_indent, filename)
        print_values("A13", 4, input_ws, output_ws, rows_indent, filename)
        print_values("B13", 5, input_ws, output_ws, rows_indent, filename)
        print_values("A14", 6, input_ws, output_ws, rows_indent, filename)
        print_values("C10", 7, input_ws, output_ws, rows_indent, filename)
        print_values("D36", 8, input_ws, output_ws, rows_indent, filename)
        print_values("G25", 9, input_ws, output_ws, rows_indent, filename)
        print_values("G22", 10, input_ws, output_ws, rows_indent, filename)
        print_values("D18", 11, input_ws, output_ws, rows_indent, filename)
        print_values("D19", 12, input_ws, output_ws, rows_indent, filename)
        print_values("B33", 13, input_ws, output_ws, rows_indent, filename)
        print_values("D16", 14, input_ws, output_ws, rows_indent, filename)
        print_values("B34", 15, input_ws, output_ws, rows_indent, filename)
        print_values("D17", 16, input_ws, output_ws, rows_indent, filename)
        print_values("E65", 17, input_ws, output_ws, rows_indent, filename)
        print_values("G65", 18, input_ws, output_ws, rows_indent, filename)
        print_values("I65", 19, input_ws, output_ws, rows_indent, filename)
        print_values("C13", 20, input_ws, output_ws, rows_indent, filename)
        print_values("G23", 21, input_ws, output_ws, rows_indent, filename)
        print_values("G24", 22, input_ws, output_ws, rows_indent, filename)

        rows_indent += 1

    try:
        if os.path.exists(output_path):
            print(f"Файл уже существует, перезаписываем: {output_path}")
        output_wb.save(output_path)
        print(f"Файл успешно сохранен! {output_path}")
    except Exception as e:
        raise fw.PluginError(f"Ошибка записи выходного файла Excel: {e}")

    return {
        "message": "Сводный отчет Бирюлево успешно сформирован!",
        "outputPath": output_path,
    }


if __name__ == "__main__":
    run()
